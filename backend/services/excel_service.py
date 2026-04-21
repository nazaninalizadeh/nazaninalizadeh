"""
Excel/CSV Import and Export Service.
Handles bulk data operations for tenants, payments, and occupancy.
"""

import io
import csv
import uuid
import logging
from datetime import datetime, timezone
from typing import Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from database import db

logger = logging.getLogger("excel_service")

HEADER_FILL = PatternFill(start_color="9F1239", end_color="9F1239", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)

# ============ Column Definitions ============

TENANT_COLUMNS = [
    ("full_name", "Nome Completo *"),
    ("codice_fiscale", "Codice Fiscale"),
    ("passport_number", "N. Passaporto *"),
    ("nationality", "Nazionalita *"),
    ("date_of_birth", "Data Nascita (YYYY-MM-DD) *"),
    ("passport_issue_date", "Rilascio Passaporto (YYYY-MM-DD) *"),
    ("passport_expiry_date", "Scadenza Passaporto (YYYY-MM-DD) *"),
    ("id_type", "Tipo Documento ID"),
    ("id_number", "Numero Documento ID"),
    ("phone", "Telefono *"),
    ("email", "Email *"),
    ("whatsapp", "WhatsApp *"),
    ("address", "Indirizzo *"),
    ("occupation", "Professione *"),
    ("notes", "Note"),
    ("deposit_amount", "Deposito"),
]

PAYMENT_COLUMNS = [
    ("tenant_email", "Email Inquilino *"),
    ("amount", "Importo *"),
    ("payment_method", "Metodo (contanti/bonifico/carta/assegno) *"),
    ("payment_date", "Data Pagamento (YYYY-MM-DD) *"),
    ("notes", "Note"),
]


def _style_header(ws, num_cols):
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                length = len(str(cell.value or ""))
                if length > max_len:
                    max_len = length
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)


# ============ Template Generation ============

def generate_tenant_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inquilini"
    for col_idx, (_, label) in enumerate(TENANT_COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=label)
    # Add sample row
    sample = [
        "Mario Rossi", "RSSMRA85M01H501Z", "AB1234567", "Italiana",
        "1985-08-01", "2020-01-15", "2030-01-15", "Carta d'identita",
        "CA12345", "+39 333 1234567", "mario@example.com", "+39 333 1234567",
        "Via Roma 1, Milano", "Ingegnere", "", "500",
    ]
    for col_idx, val in enumerate(sample, 1):
        ws.cell(row=2, column=col_idx, value=val)
    _style_header(ws, len(TENANT_COLUMNS))
    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_payment_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pagamenti"
    for col_idx, (_, label) in enumerate(PAYMENT_COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=label)
    sample = ["mario@example.com", "500.00", "bonifico", "2026-01-15", "Affitto gennaio"]
    for col_idx, val in enumerate(sample, 1):
        ws.cell(row=2, column=col_idx, value=val)
    _style_header(ws, len(PAYMENT_COLUMNS))
    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============ Export Functions ============

async def export_tenants() -> io.BytesIO:
    tenants = await db.tenants.find({}, {"_id": 0}).to_list(5000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Inquilini"

    # Headers + extra export columns
    export_cols = TENANT_COLUMNS + [
        ("total_paid", "Totale Pagato"),
        ("total_due", "Totale Dovuto"),
    ]
    for col_idx, (_, label) in enumerate(export_cols, 1):
        ws.cell(row=1, column=col_idx, value=label)
    _style_header(ws, len(export_cols))

    for row_idx, t in enumerate(tenants, 2):
        for col_idx, (key, _) in enumerate(export_cols, 1):
            val = t.get(key, "")
            ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def export_payments() -> io.BytesIO:
    payments = await db.payments.find({}, {"_id": 0}).sort("payment_date", -1).to_list(10000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pagamenti"

    headers = ["Data", "Inquilino", "Importo", "Metodo", "Note", "Creato da"]
    for col_idx, h in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header(ws, len(headers))

    for row_idx, p in enumerate(payments, 2):
        tenant = await db.tenants.find_one({"id": p.get("tenant_id", "")}, {"_id": 0})
        tenant_name = tenant.get("full_name", "") if tenant else ""
        ws.cell(row=row_idx, column=1, value=p.get("payment_date", ""))
        ws.cell(row=row_idx, column=2, value=tenant_name)
        ws.cell(row=row_idx, column=3, value=p.get("amount", 0))
        ws.cell(row=row_idx, column=4, value=p.get("payment_method", ""))
        ws.cell(row=row_idx, column=5, value=p.get("notes", ""))
        ws.cell(row=row_idx, column=6, value=p.get("created_by", ""))

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def export_occupancy() -> io.BytesIO:
    properties = await db.properties.find({}, {"_id": 0}).to_list(500)
    wb = Workbook()
    ws = wb.active
    ws.title = "Occupazione"

    headers = ["Immobile", "Codice", "Proprietario", "Stanza", "Tipo", "Stato", "Inquilino", "Affitto Mensile"]
    for col_idx, h in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header(ws, len(headers))

    row_idx = 2
    for prop in properties:
        rooms = await db.rooms.find({"property_id": prop["id"]}, {"_id": 0}).to_list(200)
        if not rooms:
            ws.cell(row=row_idx, column=1, value=prop.get("address", ""))
            ws.cell(row=row_idx, column=2, value=prop.get("property_code", ""))
            ws.cell(row=row_idx, column=3, value=prop.get("landlord_name", ""))
            ws.cell(row=row_idx, column=4, value="Nessuna stanza")
            row_idx += 1
            continue
        for room in rooms:
            tenant_name = ""
            if room.get("tenant_id"):
                tenant = await db.tenants.find_one({"id": room["tenant_id"]}, {"_id": 0})
                tenant_name = tenant.get("full_name", "") if tenant else ""
            ws.cell(row=row_idx, column=1, value=prop.get("address", ""))
            ws.cell(row=row_idx, column=2, value=prop.get("property_code", ""))
            ws.cell(row=row_idx, column=3, value=prop.get("landlord_name", ""))
            ws.cell(row=row_idx, column=4, value=room.get("room_number", ""))
            ws.cell(row=row_idx, column=5, value="Singola" if room.get("room_type") == "single" else "Doppia")
            ws.cell(row=row_idx, column=6, value="Occupata" if room.get("status") == "occupied" else "Libera")
            ws.cell(row=row_idx, column=7, value=tenant_name)
            ws.cell(row=row_idx, column=8, value=room.get("monthly_rent", 0))
            row_idx += 1

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============ Import Functions ============

async def import_tenants_from_file(file_bytes: bytes, filename: str, created_by: str) -> Tuple[int, list]:
    """
    Import tenants from Excel/CSV. Returns (success_count, errors_list).
    errors_list items: {"row": int, "error": str}
    """
    errors = []
    success = 0

    if filename.endswith(".csv"):
        rows = _parse_csv(file_bytes)
    else:
        rows = _parse_excel(file_bytes)

    if not rows:
        return 0, [{"row": 0, "error": "File vuoto o formato non valido"}]

    header = [str(c).strip().lower() for c in rows[0]]
    # Build column mapping
    field_map = _build_field_map(header, TENANT_COLUMNS)

    for row_num, row_data in enumerate(rows[1:], start=2):
        try:
            tenant = {}
            for field_key, col_idx in field_map.items():
                val = str(row_data[col_idx]).strip() if col_idx < len(row_data) and row_data[col_idx] is not None else ""
                tenant[field_key] = val

            # Validate required fields
            required = ["full_name", "passport_number", "nationality", "date_of_birth",
                        "passport_issue_date", "passport_expiry_date", "phone", "email", "whatsapp",
                        "address", "occupation"]
            missing = [f for f in required if not tenant.get(f)]
            if missing:
                errors.append({"row": row_num, "error": f"Campi obbligatori mancanti: {', '.join(missing)}"})
                continue

            # Check duplicate passport
            existing = await db.tenants.find_one({"passport_number": tenant["passport_number"]})
            if existing:
                errors.append({"row": row_num, "error": f"Passaporto duplicato: {tenant['passport_number']}"})
                continue

            # Parse deposit
            try:
                tenant["deposit_amount"] = float(tenant.get("deposit_amount", 0) or 0)
            except (ValueError, TypeError):
                tenant["deposit_amount"] = 0.0

            tenant["id"] = str(uuid.uuid4())
            tenant["total_paid"] = 0.0
            tenant["total_due"] = 0.0
            tenant["property_id"] = ""
            tenant["room_id"] = ""
            tenant["created_at"] = datetime.now(timezone.utc).isoformat()
            tenant["created_by"] = created_by

            await db.tenants.insert_one(tenant)
            success += 1

        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})

    return success, errors


async def import_payments_from_file(file_bytes: bytes, filename: str, created_by: str) -> Tuple[int, list]:
    """
    Import payments from Excel/CSV. Returns (success_count, errors_list).
    """
    errors = []
    success = 0

    if filename.endswith(".csv"):
        rows = _parse_csv(file_bytes)
    else:
        rows = _parse_excel(file_bytes)

    if not rows:
        return 0, [{"row": 0, "error": "File vuoto o formato non valido"}]

    header = [str(c).strip().lower() for c in rows[0]]
    field_map = _build_field_map(header, PAYMENT_COLUMNS)

    for row_num, row_data in enumerate(rows[1:], start=2):
        try:
            payment = {}
            for field_key, col_idx in field_map.items():
                val = str(row_data[col_idx]).strip() if col_idx < len(row_data) and row_data[col_idx] is not None else ""
                payment[field_key] = val

            tenant_email = payment.pop("tenant_email", "")
            if not tenant_email:
                errors.append({"row": row_num, "error": "Email inquilino mancante"})
                continue

            tenant = await db.tenants.find_one({"email": tenant_email.lower()})
            if not tenant:
                errors.append({"row": row_num, "error": f"Inquilino non trovato: {tenant_email}"})
                continue

            try:
                amount = float(payment.get("amount", 0))
            except (ValueError, TypeError):
                errors.append({"row": row_num, "error": "Importo non valido"})
                continue

            if amount <= 0:
                errors.append({"row": row_num, "error": "Importo deve essere positivo"})
                continue

            pay_record = {
                "id": str(uuid.uuid4()),
                "tenant_id": tenant["id"],
                "invoice_id": "",
                "amount": amount,
                "payment_method": payment.get("payment_method", ""),
                "payment_date": payment.get("payment_date", datetime.now(timezone.utc).strftime("%Y-%m-%d")),
                "notes": payment.get("notes", ""),
                "created_at": datetime.now(timezone.utc).isoformat(),
                "created_by": created_by,
            }

            await db.payments.insert_one(pay_record)
            await db.tenants.update_one({"id": tenant["id"]}, {"$inc": {"total_paid": amount}})
            success += 1

        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})

    return success, errors


# ============ Helpers ============

def _parse_csv(file_bytes: bytes) -> list:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1")
    reader = csv.reader(io.StringIO(text))
    return list(reader)


def _parse_excel(file_bytes: bytes) -> list:
    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def _build_field_map(header: list, columns: list) -> dict:
    """Map field keys to column indices, matching by header label or key name."""
    field_map = {}
    for field_key, label in columns:
        label_lower = label.lower().replace("*", "").strip()
        key_lower = field_key.lower()
        for idx, h in enumerate(header):
            h_clean = h.lower().replace("*", "").strip()
            if h_clean == label_lower or h_clean == key_lower or h_clean.startswith(key_lower):
                field_map[field_key] = idx
                break
    return field_map
